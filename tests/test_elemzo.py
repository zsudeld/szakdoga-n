"""
tests/test_elemzo.py
====================
Egységtesztek a hibrid sentiment elemző alapfüggvényeihez.

Futtatás: pytest tests/
"""

import pytest


# ---------------------------------------------------------------------------
# _hibrid_kategoria
# ---------------------------------------------------------------------------

from huspacy_elemzo import _hibrid_kategoria


class TestHibridKategoria:

    def test_egyezes_pozitiv(self):
        kat, _ = _hibrid_kategoria(0.5, "pozitív", 0.8)
        assert kat == "pozitív"

    def test_egyezes_negativ(self):
        kat, _ = _hibrid_kategoria(-0.5, "negatív", 0.8)
        assert kat == "negatív"

    def test_egyezes_semleges(self):
        kat, _ = _hibrid_kategoria(0.0, "semleges", 0.9)
        assert kat == "semleges"

    def test_nytk_magas_konfidencia_dont(self):
        """NYTK magas konfidenciával dönt HuSpaCy gyenge jelével szemben."""
        kat, _ = _hibrid_kategoria(0.05, "negatív", 0.80)
        assert kat == "negatív"

    def test_huspacy_eros_jel_dont(self):
        """Erős szótári jel dönt mérsékelt NYTK-val szemben."""
        kat, _ = _hibrid_kategoria(0.40, "semleges", 0.52)
        assert kat == "pozitív"

    def test_gyenge_jelek_semleges(self):
        """Mindkét jel tényleg gyenge → semleges. mc=0.49 < 0.60, abs(hp)=0.05 < 0.22"""
        kat, _ = _hibrid_kategoria(0.05, "pozitív", 0.49)
        assert kat == "semleges"

    def test_visszateres_tuple(self):
        eredmeny = _hibrid_kategoria(0.3, "pozitív", 0.7)
        assert isinstance(eredmeny, tuple)
        assert len(eredmeny) == 2
        assert isinstance(eredmeny[1], str)

    def test_megalapozas_nem_ures(self):
        _, megalapozas = _hibrid_kategoria(0.3, "pozitív", 0.7)
        assert len(megalapozas) > 0


# ---------------------------------------------------------------------------
# _feldolgoz_nytk_eredmeny
# ---------------------------------------------------------------------------

from huspacy_elemzo import _feldolgoz_nytk_eredmeny


class TestFeldolgozNytkEredmeny:

    def test_none_input(self):
        """A hiányzó modellválasz hiba, nem kitalált semleges."""
        kat, conf = _feldolgoz_nytk_eredmeny(None)
        assert kat == "hiba"
        assert conf == 0.0

    def test_ures_lista(self):
        kat, conf = _feldolgoz_nytk_eredmeny([])
        assert kat == "hiba"
        assert conf == 0.0

    def test_ures_valaszt_ado_pipeline(self):
        """A None-t vagy üres listát visszaadó helyettesítő pipeline hibát ad."""
        class _UresPipeline:
            def __call__(self, szoveg):
                return None
        kat, conf = _modell_elemez("teszt", _UresPipeline())
        assert (kat, conf) == ("hiba", 0.0)

    def test_kivetelt_do_pipeline(self):
        class _HibasPipeline:
            def __call__(self, szoveg):
                raise RuntimeError("modell hiba")
        kat, conf = _modell_elemez("teszt", _HibasPipeline())
        assert (kat, conf) == ("hiba", 0.0)

    def test_harom_osztalyos_pozitiv(self):
        raw = [
            {"label": "LABEL_0", "score": 0.05},
            {"label": "LABEL_1", "score": 0.10},
            {"label": "LABEL_2", "score": 0.85},
        ]
        kat, conf = _feldolgoz_nytk_eredmeny(raw)
        assert kat == "pozitív"
        assert conf == pytest.approx(0.85)

    def test_harom_osztalyos_negativ(self):
        raw = [
            {"label": "LABEL_0", "score": 0.90},
            {"label": "LABEL_1", "score": 0.05},
            {"label": "LABEL_2", "score": 0.05},
        ]
        kat, conf = _feldolgoz_nytk_eredmeny(raw)
        assert kat == "negatív"
        assert conf == pytest.approx(0.90)

    def test_harom_osztalyos_semleges(self):
        raw = [
            {"label": "LABEL_0", "score": 0.10},
            {"label": "LABEL_1", "score": 0.80},
            {"label": "LABEL_2", "score": 0.10},
        ]
        kat, conf = _feldolgoz_nytk_eredmeny(raw)
        assert kat == "semleges"

    def test_ot_osztalyos_pozitiv(self):
        raw = [
            {"label": "LABEL_0", "score": 0.02},
            {"label": "LABEL_1", "score": 0.03},
            {"label": "LABEL_2", "score": 0.05},
            {"label": "LABEL_3", "score": 0.30},
            {"label": "LABEL_4", "score": 0.60},
        ]
        kat, conf = _feldolgoz_nytk_eredmeny(raw)
        assert kat == "pozitív"
        assert conf == pytest.approx(0.90)

    def test_explicit_positive_label(self):
        raw = [{"label": "POSITIVE", "score": 0.92}]
        kat, conf = _feldolgoz_nytk_eredmeny(raw)
        assert kat == "pozitív"
        assert conf == pytest.approx(0.92)

    def test_explicit_negative_label(self):
        raw = [{"label": "NEGATIVE", "score": 0.88}]
        kat, conf = _feldolgoz_nytk_eredmeny(raw)
        assert kat == "negatív"

    def test_dict_input_wrapped(self):
        """Egyetlen dict-et listába kell csomagolni."""
        raw = {"label": "LABEL_2", "score": 0.75}
        kat, _ = _feldolgoz_nytk_eredmeny(raw)
        assert kat == "pozitív"


# ---------------------------------------------------------------------------
# _szotari_pontszam – SpaCy nélkül, mock doc-cal tesztelve
# ---------------------------------------------------------------------------

from huspacy_elemzo import (
    _szotari_pontszam,
    _token_talal_pozitiv,
    _token_talal_negativ,
)


class MockToken:
    """Minimális SpaCy token mock a szótári függvények teszteléséhez."""

    def __init__(self, text: str, lemma: str = "", pos: str = "NOUN",
                 is_stop: bool = False, is_space: bool = False,
                 is_punct: bool = False, children=None, head_lemma: str = ""):
        self.text = text
        self.lemma_ = lemma or text
        self.pos_ = pos
        self.is_stop = is_stop
        self.is_space = is_space
        self.is_punct = is_punct
        self.children = children or []
        self._head_lemma = head_lemma
        self.i = 0

    @property
    def head(self):
        class _Head:
            def __init__(self, lemma):
                self.lemma_ = lemma
                self.text = lemma
                self.dep_ = "ROOT"
        return _Head(self._head_lemma)

    @property
    def sent(self):
        return self


class MockDoc:
    def __init__(self, tokens):
        self._tokens = tokens
        for i, t in enumerate(tokens):
            t.i = i

    def __iter__(self):
        return iter(self._tokens)

    def __getitem__(self, idx):
        return self._tokens[idx]


class TestTokenTalal:

    def test_pozitiv_szoalak(self):
        tok = MockToken("jó")
        assert _token_talal_pozitiv(tok)

    def test_pozitiv_lemma(self):
        tok = MockToken("jók", lemma="jó")
        assert _token_talal_pozitiv(tok)

    def test_negativ_szoalak(self):
        tok = MockToken("rossz")
        assert _token_talal_negativ(tok)

    def test_semleges_szo(self):
        tok = MockToken("asztal")
        assert not _token_talal_pozitiv(tok)
        assert not _token_talal_negativ(tok)

    def test_tobbszavas_nem_illeszkedik(self):
        """Többszavas kifejezések (ha maradt volna) soha nem illeszkednek."""
        tok = MockToken("nem ajánlom")
        assert not _token_talal_negativ(tok)


class TestSzotariPontszam:

    def test_pozitiv_szoveg_pozitiv_pontot_ad(self):
        tokens = [
            MockToken("kiváló", pos="ADJ"),
            MockToken("remek", pos="ADJ"),
            MockToken("jó", pos="ADJ"),
        ]
        doc = MockDoc(tokens)
        pont = _szotari_pontszam(doc)
        assert pont > 0

    def test_negativ_szoveg_negativ_pontot_ad(self):
        tokens = [
            MockToken("rossz", pos="ADJ"),
            MockToken("borzasztó", pos="ADJ"),
        ]
        doc = MockDoc(tokens)
        pont = _szotari_pontszam(doc)
        assert pont < 0

    def test_ures_doc_nulla(self):
        doc = MockDoc([])
        assert _szotari_pontszam(doc) == 0.0

    def test_eredmeny_minus1_es_1_kozott(self):
        tokens = [MockToken("kiváló", pos="ADJ")] * 50
        doc = MockDoc(tokens)
        pont = _szotari_pontszam(doc)
        assert -1.0 <= pont <= 1.0

    def test_csak_stopszavak_nulla(self):
        tokens = [
            MockToken("a", is_stop=True, pos="DET"),
            MockToken("az", is_stop=True, pos="DET"),
        ]
        doc = MockDoc(tokens)
        assert _szotari_pontszam(doc) == 0.0


# ---------------------------------------------------------------------------
# generalj_temakat
# ---------------------------------------------------------------------------

from huspacy_elemzo import generalj_temakat


class TestGeneraljTemakat:

    def test_ures_lista(self):
        assert generalj_temakat([]) == []

    def test_visszater_lista(self):
        szovegek = ["Fantasztikus kiszolgálás volt", "Rossz étel hideg volt"]
        eredmeny = generalj_temakat(szovegek)
        assert isinstance(eredmeny, list)
        assert len(eredmeny) == len(szovegek)

    def test_minden_sor_kap_kulcsszavat(self):
        szovegek = ["alma körte barack", "autó motor kerék", "ház ablak ajtó"]
        eredmeny = generalj_temakat(szovegek)
        for e in eredmeny:
            assert isinstance(e, str)
            assert len(e) > 0


# ---------------------------------------------------------------------------
# intelligens_csv_beolvasas
# ---------------------------------------------------------------------------

import pandas as pd
import tempfile
import os

from utils import intelligens_csv_beolvasas, optimalis_oszlop


class TestCsvBeolvasas:

    def setup_method(self):
        self._temp_files: list[str] = []

    def _tmp_csv(self, tartalom: str, encoding: str = "utf-8") -> str:
        f = tempfile.NamedTemporaryFile(
            mode="w", suffix=".csv", encoding=encoding,
            delete=False, newline=""
        )
        f.write(tartalom)
        f.close()
        self._temp_files.append(f.name)
        return f.name

    def teardown_method(self):
        for path in self._temp_files:
            try:
                os.chmod(path, 0o644)
                os.unlink(path)
            except OSError:
                pass
        self._temp_files.clear()

    def test_pontosvesszo_elvalaszto(self):
        path = self._tmp_csv("nev;velemeny\nKati;Nagyon jo volt\nPeti;Rossz volt\n")
        df = intelligens_csv_beolvasas(path)
        assert df is not None
        assert "velemeny" in df.columns

    def test_vesszos_elvalaszto(self):
        path = self._tmp_csv("name,review\nAnn,Great service\nBob,Terrible food\n")
        df = intelligens_csv_beolvasas(path)
        assert df is not None
        assert len(df) == 2
        # regresszió: a vesszős fájl ne egyetlen oszlopként érkezzen
        assert list(df.columns) == ["name", "review"]

    def test_vesszo_a_szovegben_pontosvesszos_fajl(self):
        path = self._tmp_csv('nev;velemeny\nKati;"Jó volt, nagyon"\nPeti;Rossz volt\n')
        df = intelligens_csv_beolvasas(path)
        assert list(df.columns) == ["nev", "velemeny"]
        assert df["velemeny"].iloc[0] == "Jó volt, nagyon"

    def test_egyoszlopos_fajl_idezett_vesszovel(self):
        path = self._tmp_csv('velemeny\n"Jó volt, nagyon"\n"Rossz, de hasznos"\n')
        df = intelligens_csv_beolvasas(path)
        assert list(df.columns) == ["velemeny"]
        assert len(df) == 2

    def test_tabulatoros_fajl(self):
        path = self._tmp_csv("nev\tvelemeny\nKati\tJo volt\nPeti\tRossz volt\n")
        df = intelligens_csv_beolvasas(path)
        assert list(df.columns) == ["nev", "velemeny"]

    def test_nem_letezo_fajl_none(self):
        df = intelligens_csv_beolvasas("/nem/letezik/fajl.csv")
        assert df is None

    def test_permission_error(self):
        if os.getuid() == 0:
            pytest.skip("Root felhasználóként a chmod 000 nem okoz PermissionError-t")
        path = self._tmp_csv("a;b\n1;2\n")
        os.chmod(path, 0o000)
        with pytest.raises(PermissionError):
            intelligens_csv_beolvasas(path)


class TestOptimalisOszlop:

    def test_leghosszabb_szoveg_oszlop(self):
        df = pd.DataFrame({
            "id": ["1", "2", "3"],
            "szoveg": [
                "Ez egy hosszabb véleményes szöveg",
                "Ez is egy hosszabb szöveg lesz",
                "Szintén hosszabb vélemény van itt",
            ],
        })
        oszlop = optimalis_oszlop(df)
        assert oszlop == "szoveg"

    def test_egyszeru_fallback(self):
        df = pd.DataFrame({"a": ["x", "y"]})
        oszlop = optimalis_oszlop(df)
        assert oszlop == "a"


# ---------------------------------------------------------------------------
# _negalt_e: regresszió a mondatszeletek (Span) abszolút indexelésére
# ---------------------------------------------------------------------------

spacy = pytest.importorskip("spacy")
from spacy.tokens import Doc  # noqa: E402

import huspacy_elemzo as _he  # noqa: E402


def _ketmondatos_doc(masodik_negalt: bool):
    """
    Két mondatos, modell nélkül összeállított spaCy Doc.
    1. mondat: "Ez jó ."
    2. mondat: "Ez nem volt jó ."  (negált)  vagy  "Ez volt nagyon jó ." (nem negált)
    """
    nlp = spacy.blank("hu")
    if masodik_negalt:
        words  = ["Ez", "jó", ".", "Ez", "nem", "volt", "jó", "."]
        lemmas = ["ez", "jó", ".", "ez", "nem", "van", "jó", "."]
        heads  = [1, 1, 1, 5, 5, 5, 5, 5]
        deps   = ["nsubj", "ROOT", "punct", "nsubj", "advmod", "ROOT", "xcomp", "punct"]
        pos    = ["PRON", "ADJ", "PUNCT", "PRON", "PART", "AUX", "ADJ", "PUNCT"]
    else:
        words  = ["Ez", "jó", ".", "Ez", "volt", "nagyon", "jó", "."]
        lemmas = ["ez", "jó", ".", "ez", "van", "nagyon", "jó", "."]
        heads  = [1, 1, 1, 4, 4, 6, 4, 4]
        deps   = ["nsubj", "ROOT", "punct", "nsubj", "ROOT", "advmod", "xcomp", "punct"]
        pos    = ["PRON", "ADJ", "PUNCT", "PRON", "AUX", "ADV", "ADJ", "PUNCT"]
    starts = [True, False, False, True, False, False, False, False]
    return Doc(nlp.vocab, words=words, spaces=[True] * 8, lemmas=lemmas, heads=heads,
               deps=deps, sent_starts=starts, pos=pos)


class TestNegaltESpan:

    def test_nem_negalt_szo_kesobbi_mondatban_nem_dob_hibat(self):
        doc = _ketmondatos_doc(masodik_negalt=False)
        mondat = list(doc.sents)[1]
        token = [t for t in mondat if t.text == "jó"][0]
        assert token.i == 6 and len(mondat) == 5      # abszolút index > a szelet hossza
        assert _he._negalt_e(token, mondat) is False

    def test_negalt_szo_kesobbi_mondatban(self):
        doc = _ketmondatos_doc(masodik_negalt=True)
        mondat = list(doc.sents)[1]
        token = [t for t in mondat if t.text == "jó"][0]
        assert _he._negalt_e(token, mondat) is True

    def test_mondatszintu_bontas_minden_mondatra_ertelmes(self):
        doc = _ketmondatos_doc(masodik_negalt=False)
        kimenet = _he._mondatszintu_osszefoglalo_doc(doc)
        assert "M1:poz" in kimenet
        assert "M2:poz" in kimenet      # a hiba idején itt "M2:sem(+0.00)" állt

    def test_negalt_mondat_negativ_iranyba_mozdul(self):
        doc = _ketmondatos_doc(masodik_negalt=True)
        mondat = list(doc.sents)[1]
        assert _he._szotari_pontszam(mondat) < _he._szotari_pontszam(list(doc.sents)[0])


# ---------------------------------------------------------------------------
# kiertekeles.py
# ---------------------------------------------------------------------------

from kiertekeles import (  # noqa: E402
    CIMKEK, cimke_normalizal, felosztas, kappa, konfuzios_matrix, ertekel, metrikak,
)


class TestCimkeNormalizal:

    def test_alias_ertekek(self):
        assert cimke_normalizal("Pozitív") == "pozitív"
        assert cimke_normalizal(" negativ ") == "negatív"
        assert cimke_normalizal("NEUTRAL") == "semleges"

    def test_hianyzo_ertek_none(self):
        assert cimke_normalizal(None) is None
        assert cimke_normalizal(float("nan")) is None
        assert cimke_normalizal("  ") is None

    def test_ismeretlen_cimke_hibat_dob(self):
        with pytest.raises(ValueError):
            cimke_normalizal("talán")

    def test_numerikus_cimke_nem_talalgat(self):
        with pytest.raises(ValueError):
            cimke_normalizal("1")


class TestMetrikak:

    def test_kezzel_szamolt_ertekek(self):
        gold    = ["pozitív"] * 4 + ["semleges", "negatív"]
        becsles = ["pozitív", "pozitív", "pozitív", "semleges", "semleges", "pozitív"]
        m = metrikak(gold, becsles)
        assert m["pontossag"] == pytest.approx(4 / 6)
        assert m["osztaly"]["pozitív"]["f1"] == pytest.approx(0.75)
        assert m["osztaly"]["semleges"]["f1"] == pytest.approx(2 / 3)
        assert m["osztaly"]["negatív"]["f1"] == 0.0
        assert m["makro_f1"] == pytest.approx((0.75 + 2 / 3 + 0.0) / 3)

    def test_hianyzo_osztaly_nem_torzitja_a_makro_f1_et(self):
        m = metrikak(["pozitív", "pozitív"], ["pozitív", "pozitív"])
        assert m["makro_f1"] == pytest.approx(1.0)

    def test_ervenytelen_becsles_teves(self):
        m = metrikak(["pozitív", "negatív"], ["pozitív", "hiba"])
        assert m["pontossag"] == pytest.approx(0.5)
        assert m["ervenytelen_becslesek"] == 1

    def test_hossz_eltereseknel_hiba(self):
        with pytest.raises(ValueError):
            metrikak(["pozitív"], [])

    def test_konfuzios_matrix_sorai_es_oszlopai(self):
        gold    = ["pozitív", "pozitív", "negatív"]
        becsles = ["pozitív", "semleges", "negatív"]
        df = konfuzios_matrix(gold, becsles)
        assert list(df.index) == CIMKEK
        assert df.loc["pozitív", "semleges"] == 1
        assert df.loc["negatív", "negatív"] == 1
        assert "hiba" not in df.columns

    def test_kappa_kezzel_szamolt(self):
        g1 = ["pozitív", "pozitív", "negatív", "negatív"]
        g2 = ["pozitív", "pozitív", "negatív", "pozitív"]
        assert kappa(g1, g2) == pytest.approx(0.5)
        assert kappa(g1, g1) == pytest.approx(1.0)


class TestFelosztas:

    def test_reprodukalhato_es_diszjunkt(self):
        gold = ["pozitív"] * 10 + ["negatív"] * 10 + ["semleges"] * 10
        dev1, teszt1 = felosztas(gold, 0.7, seed=1)
        dev2, teszt2 = felosztas(gold, 0.7, seed=1)
        assert (dev1, teszt1) == (dev2, teszt2)
        assert not set(dev1) & set(teszt1)
        assert sorted(dev1 + teszt1) == list(range(30))
        assert len(teszt1) == 21

    def test_retegzett_aranyok(self):
        gold = ["pozitív"] * 10 + ["negatív"] * 10 + ["semleges"] * 10
        _, teszt = felosztas(gold, 0.7, seed=3)
        assert {gold[i] for i in teszt} == set(CIMKEK)

    def test_ritka_osztalynal_is_lefut(self):
        gold = ["pozitív"] * 8 + ["negatív"]          # egyetlen negatív
        dev, teszt = felosztas(gold, 0.5, seed=0)
        assert len(dev) + len(teszt) == 9

    def test_ervenytelen_arany(self):
        with pytest.raises(ValueError):
            felosztas(["pozitív", "negatív"], 1.5)


class _HelyettesitoMotor:
    """Modell nélküli elemző: a szöveg alapján előre megadott kategóriákat ad."""

    def __init__(self, tabla):
        self.tabla = tabla

    def elemzes_batch(self, szovegek):
        for sz in szovegek:
            szotari, hubert, hibrid = self.tabla[sz]
            yield {"kategoria": szotari, "hunbert_kategoria": hubert, "hunbert_confidence": 0.9,
                   "pontszam": 0.1, "hibrid_kategoria": hibrid, "hibrid_megalapozas": "teszt"}


class TestErtekel:

    def _adat(self):
        szovegek = [f"szoveg {i}" for i in range(20)]
        gold = (["pozitív"] * 8) + (["negatív"] * 8) + (["semleges"] * 4)
        # a hibrid mindig helyes, a szótári mindig "semleges", a HuBERT mindig helyes
        tabla = {sz: ("semleges", g, g) for sz, g in zip(szovegek, gold)}
        df = pd.DataFrame({"szoveg": szovegek, "cimke": gold})
        return df, tabla

    def test_teljes_lefutas_es_szerkezet(self):
        df, tabla = self._adat()
        er = ertekel(df, "szoveg", "cimke", _HelyettesitoMotor(tabla), teszt_arany=0.5, seed=1)
        o = er["osszehasonlitas"]
        assert set(o["Részminta"]) == {"fejlesztői", "teszt", "összes"}
        teszt = o[(o["Részminta"] == "teszt") & (o["Konfiguráció"] == "C) Hibrid módszer")].iloc[0]
        assert teszt["Pontosság"] == 1.0 and teszt["Makro-F1"] == 1.0
        szotari = o[(o["Részminta"] == "összes") & (o["Konfiguráció"] == "A) Szótári módszer")].iloc[0]
        assert szotari["Pontosság"] == pytest.approx(4 / 20)
        assert er["hibak"].empty
        assert er["meta"]["Címkézett szövegek száma"] == 20

    def test_hianyzo_cimke_sorok_kimaradnak(self):
        df, tabla = self._adat()
        df.loc[0, "cimke"] = None
        er = ertekel(df, "szoveg", "cimke", _HelyettesitoMotor(tabla), teszt_arany=0.5)
        assert er["meta"]["Címkézett szövegek száma"] == 19
        assert er["meta"]["Kihagyott sorok (hiányzó címke vagy szöveg)"] == 1

    def test_hibas_besorolasok_listaja(self):
        df, tabla = self._adat()
        tabla["szoveg 0"] = ("pozitív", "pozitív", "negatív")     # a hibrid téved
        er = ertekel(df, "szoveg", "cimke", _HelyettesitoMotor(tabla), teszt_arany=0.5, seed=0)
        osszes = er["osszehasonlitas"]
        hibrid_osszes = osszes[(osszes["Részminta"] == "összes")
                               & (osszes["Konfiguráció"] == "C) Hibrid módszer")].iloc[0]
        assert hibrid_osszes["Pontosság"] == pytest.approx(19 / 20)
        # a hibás besorolás pontosan akkor szerepel a listában, ha a teszt részre esett
        _, teszt_idx = felosztas(list(df["cimke"]), 0.5, seed=0)
        assert (0 in teszt_idx) == (len(er["hibak"]) == 1)
        if len(er["hibak"]) == 1:
            assert er["hibak"].iloc[0]["Szöveg"] == "szoveg 0"
            assert er["hibak"].iloc[0]["Hibrid"] == "negatív"

    def test_kappa_a_masodik_cimkezovel(self):
        df, tabla = self._adat()
        df["cimke2"] = df["cimke"]
        er = ertekel(df, "szoveg", "cimke", _HelyettesitoMotor(tabla), gold2_oszlop="cimke2")
        assert er["kappa"] == pytest.approx(1.0)

    def test_hianyzo_oszlop_hibat_dob(self):
        df, tabla = self._adat()
        with pytest.raises(KeyError):
            ertekel(df, "nincs_ilyen", "cimke", _HelyettesitoMotor(tabla))

# ---------------------------------------------------------------------------
# Regressziós tesztek a javításokra (hiba-állapot, képletvédelem,
# soreldobás-jelzés, BERTopic-forrásjelölés)
# ---------------------------------------------------------------------------

from huspacy_elemzo import _hiba_eredmeny, _modell_elemez  # noqa: E402


class TestHibaEredmeny:
    """A feldolgozási hiba ne nézzen ki semleges (érvényes) becslésnek."""

    def test_hiba_eredmeny_hiba_allapot(self):
        e = _hiba_eredmeny("tesztszoveg", "feldolgozasi hiba")
        assert e["hibrid_kategoria"] == "hiba"
        assert e["hunbert_kategoria"] == "hiba"

    def test_hiba_ervenytelen_becsles_a_metrikakban(self):
        e = _hiba_eredmeny("tesztszoveg", "hiba")
        m = metrikak(["semleges"], [e["hibrid_kategoria"]])
        assert m["ervenytelen_becslesek"] == 1
        assert m["pontossag"] == 0.0

    def test_modell_hiba_nem_semleges(self):
        kat, conf = _modell_elemez("szoveg", None)
        assert kat == "hiba"
        assert conf == 0.0


class TestExcelFormulaVedelem:
    """A '=' jellel kezdődő bemeneti szöveg ne kerüljön képletként az Excelbe."""

    def _df(self):
        return pd.DataFrame({
            "eredeti_szoveg": ["=1+1", "=HYPERLINK(\"http://rosszold.hu\")", "rendben"],
            "hibrid_kategoria": ["pozitív", "pozitív", "semleges"],
        })

    def test_riport_szovegkent_tarol(self, tmp_path):
        from riport_generator import generalj_riportot
        import openpyxl
        ut = tmp_path / "formula_riport.xlsx"
        generalj_riportot(self._df(), {
            "Pozitív (Hibrid)": 2, "Negatív (Hibrid)": 0,
            "Semleges (Hibrid)": 1, "Hiba (feldolgozás)": 0,
            "Átlagos HuSpaCy pont": 0.0,
        }, kimenet=str(ut))
        wb = openpyxl.load_workbook(str(ut))
        ws = wb["Elemzési Eredmények"]
        c1, c2, c3 = ws["B2"], ws["B3"], ws["B4"]
        assert c1.data_type == "s" and c1.value == "=1+1"
        assert c2.data_type == "s" and c2.value.startswith("=HYPERLINK")
        assert c3.data_type == "s" and c3.value == "rendben"

    def test_kiertekeles_export_szovegkent_tarol(self, tmp_path):
        from kiertekeles import mentes_excel
        import openpyxl
        er = {
            "osszehasonlitas": pd.DataFrame({"Konfiguráció": ["teszt"], "Pontosság": [1.0]}),
            "konfuzio": pd.DataFrame([[1, 0], [0, 1]]),
            "hibak": pd.DataFrame([{"Szöveg": "=1+1", "Valós": "pozitív"}]),
            "meta": {"Szöveg": "=2+2"},
            "kappa": None,
        }
        ut = tmp_path / "kiertekeles.xlsx"
        mentes_excel(er, str(ut))
        wb = openpyxl.load_workbook(str(ut))
        talalt = [
            c for lap in wb.worksheets for sor in lap.iter_rows() for c in sor
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        assert talalt, "az exportált képletszerű szövegnek szerepelnie kell"
        for c in talalt:
            assert c.data_type == "s"


class TestHibasSorokJelzese:
    """A többletmezős sorok kihagyása ne legyen csendes."""

    def test_tobbletmezos_sor_szerkezeti_hiba(self, tmp_path):
        """A fejlécnél több mezős rekord után a beolvasó világos hibával áll le –
        a csendes kihagyás/index-értelmezés csúszott adatokat elemelne."""
        ut = tmp_path / "hibas_sor.csv"
        ut.write_text(
            "nev;velemeny\nKati;jo\nPeti;rossz;FELESLEGES\nAnna;jo\n",
            encoding="utf-8",
        )
        with pytest.raises(ValueError, match="szerkezete hibás.*3. rekord"):
            intelligens_csv_beolvasas(str(ut))

    def test_tobbletmezos_elso_adatsor_is_szerkezeti_hiba(self, tmp_path):
        ut = tmp_path / "hibas_elso.csv"
        ut.write_text(
            "nev;velemeny\nPeti;rossz;FELESLEGES\nAnna;jo\n",
            encoding="utf-8",
        )
        with pytest.raises(ValueError, match="szerkezete hibás"):
            intelligens_csv_beolvasas(str(ut))

    def test_ures_sorok_nem_szerkezeti_hiba(self, tmp_path):
        """A szabályos üres rekordok nem hibásak – a beolvasás rendben lefut."""
        ut = tmp_path / "ures_sorok.csv"
        ut.write_text(
            "nev;velemeny\nKati;jo\n\nAnna;jo\n",
            encoding="utf-8",
        )
        df = intelligens_csv_beolvasas(str(ut))
        assert df is not None
        assert len(df) >= 2

    def test_ideszett_sortores_nem_hibas_sor(self, tmp_path, capsys):
        ut = tmp_path / "idezet.csv"
        ut.write_text(
            'nev;velemeny\nKati;"tobb\nsoros, gondolat"\nAnna;jo\n',
            encoding="utf-8",
        )
        df = intelligens_csv_beolvasas(str(ut))
        kimenet = capsys.readouterr().out
        assert df is not None and len(df) == 2
        assert "FIGYELEM" not in kimenet

    def test_ures_fajl_tisztan_kezelve(self, tmp_path, capsys):
        ut = tmp_path / "ures.csv"
        ut.write_text("", encoding="utf-8")
        df = intelligens_csv_beolvasas(str(ut))
        kimenet = capsys.readouterr().out
        assert df is None
        assert "ures" in kimenet
        assert "Hiba (" not in kimenet  # nem zajos tízkísérletes kimenet


class TestBertopicForrasJeloles:
    """A Témaelemzés lap csak akkor ígérjen BERTopicot, ha tényleg futott."""

    def _lap(self, df):
        from openpyxl import Workbook
        from riport_generator import _lap_temak
        wb = Workbook()
        _lap_temak(wb, df)
        return wb["Témaelemzés"]

    def test_bertopic_futott(self):
        df = pd.DataFrame({
            "bertopic_tema": ["Téma 1", "Téma 1", "Téma 2"],
            "tema_kulcsszavak": ["a", "b", "c"],
        })
        ws = self._lap(df)
        assert "BERTOPIC" in ws["A1"].value

    def test_bertopic_nem_futott(self):
        df = pd.DataFrame({
            "tema_kulcsszavak": ["a", "b", "c"],
        })
        ws = self._lap(df)
        assert "TF-IDF" in ws["A1"].value
        assert "BERTopic NEM FUTOTT" in ws["A1"].value
        # a megjegyzés is jelezze
        assert "NEM BERTopic" in ws["A2"].value


class TestHibasNlpSorBatchben:
    """A hibázó HuSpaCy-motor egy sorra: a sor 'hiba' állapotot kap, a többi
    sor feldolgozása megmarad, és a hibás sor nem lesz érvényes semleges találat."""

    def _motor(self, hibas_szoveg):
        from huspacy_elemzo import FejlettSentimentElemzo
        motor = FejlettSentimentElemzo.__new__(FejlettSentimentElemzo)
        motor.nlp = _HibasNlp(hibas_szoveg)
        motor.sentiment = None
        return motor

    def test_hibas_sor_hiba_allapot(self):
        motor = self._motor("hibas szoveg")
        szovegek = ["nagyon jó", "hibas szoveg", "rossz"]
        eredmenyek = list(motor.elemzes_batch(szovegek))
        assert len(eredmenyek) == 3
        # a hibás sor hiba-állapotot kap, nem semlegest
        assert eredmenyek[1]["hibrid_kategoria"] == "hiba"
        assert eredmenyek[1]["kategoria"] == "hiba"
        # a többi sor feldolgozása megmarad
        assert eredmenyek[0]["hibrid_kategoria"] == "pozitív"
        assert eredmenyek[2]["hibrid_kategoria"] == "negatív"
        # a hibás sor a kiértékelésben érvénytelen becslés
        m = metrikak(["pozitív"], [eredmenyek[1]["hibrid_kategoria"]])
        assert m["ervenytelen_becslesek"] == 1
        assert m["pontossag"] == 0.0


class TestStatisztikaHibaSor:
    """A feldolgozási hibák jelenjenek meg a Statisztikai Összefoglaló lapban."""

    def test_hiba_kategoria_latszik(self):
        from openpyxl import Workbook
        from riport_generator import _lap_statisztika
        df = pd.DataFrame({"hunbert_confidence": [0.5, 0.6, 0.4]})
        stats = {
            "Pozitív (Hibrid)": 1, "Negatív (Hibrid)": 1,
            "Semleges (Hibrid)": 2, "Hiba (feldolgozás)": 1,
            "Átlagos HuSpaCy pont": 0.1,
        }
        wb = Workbook()
        _lap_statisztika(wb, df, stats)
        ws = wb["Statisztikai Összefoglaló"]
        # a 4 kategória sorai: 5..8
        assert ws["A8"].value == "Hiba – feldolgozási hiba"
        assert ws["B8"].value == 1
        # a százalékok a négy kategória összegéből számolódnak
        assert ws["C5"].value == pytest.approx(1 / 5 * 100)


# ---------------------------------------------------------------------------
# Mock NLP-motor a batch-integrációs teszthez (modellletöltés nélkül)
# ---------------------------------------------------------------------------

class _MockNlpToken(MockToken):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.dep_ = "nsubj"


class _MockNlpDoc(MockDoc):
    @property
    def sents(self):
        return []

    @property
    def ents(self):
        return []


class _HibasNlp:
    """A megadott szövegre hibázó, egyébként mock dokumentumot adó NLP-motor."""

    def __init__(self, hibas_szoveg):
        self.hibas = hibas_szoveg

    def pipe(self, szovegek, batch_size=32):
        for s in szovegek:
            if s == self.hibas:
                raise RuntimeError("mock HuSpaCy hiba")
            yield self._doc(s)

    def __call__(self, szoveg):
        if szoveg == self.hibas:
            raise RuntimeError("mock HuSpaCy hiba")
        return self._doc(szoveg)

    def _doc(self, szoveg):
        return _MockNlpDoc([_MockNlpToken(w, pos="ADJ") for w in szoveg.split()])
